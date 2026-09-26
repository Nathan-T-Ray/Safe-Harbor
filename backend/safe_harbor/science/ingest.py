"""Reproducible, source-backed compact ingestion. Run with python -m ...ingest.

No biological conclusions are copied from the publication. The workbook inventory
preserves original rows, including blank/merged-cell structure, separately from
the specifically normalized tables used by tools.
"""
from __future__ import annotations

import gzip
import hashlib
import json
import re
import time
from pathlib import Path
from urllib.request import urlopen

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "data" / "safe_harbor"
RAW = DATA / "raw"
NORMALIZED = DATA / "normalized"
STUDY = "https://elifesciences.org/articles/79592"
VERSION = "autio-79592-v2-gencode36-comp-pseudo-2026-09-26"
NAMES = {"Pansio": ("pansio-1", "Pansio-1", "chr1"),
         "Olônne": ("olonne-18", "Olônne-18", "chr18"),
         "Keppel": ("keppel-19", "Keppel-19", "chr19")}


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def write_json(name: str, value) -> None:
    NORMALIZED.mkdir(parents=True, exist_ok=True)
    (NORMALIZED / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def download(url: str, name: str) -> dict:
    path = RAW / name
    if not path.exists():
        path.write_bytes(urlopen(url, timeout=120).read())
    raw = path.read_bytes()
    return {"source_id": name, "url": url, "path": str(path.relative_to(ROOT)),
            "sha256": digest(raw), "bytes": len(raw)}


def workbooks() -> tuple[list[dict], list[dict], dict]:
    sources, inventories, expression = [], [], {}
    for number in (1, 3, 4, 5, 6):
        name = f"elife-79592-supp{number}.xlsx"
        source = download(f"https://cdn.elifesciences.org/articles/79592/elife-79592-supp{number}-v2.xlsx", name)
        source.update(license="CC BY 4.0", attribution="Autio et al., eLife 2024;13:e79592", source_version="v2")
        sources.append(source)
        book = openpyxl.load_workbook(RAW / name, data_only=True)
        sheets = []
        for sheet in book:
            rows = [{"source_row": index, "values": list(values)}
                    for index, values in enumerate(sheet.iter_rows(values_only=True), 1)]
            sheets.append({"sheet": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column,
                           "merged_ranges": [str(x) for x in sheet.merged_cells.ranges], "rows": rows})
        inventories.append({"source_id": name, "sha256": source["sha256"], "sheets": sheets})
        if number in (5, 6):
            context = "H1" if number == 5 else "H9"
            expression[context] = {}
            for sheet in book.worksheets[1:5]:
                label = next((x[0] for prefix, x in NAMES.items() if sheet.title.startswith(prefix)), "untargeted")
                headers = [c.value for c in sheet[1]]
                records = []
                for rowno, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    if not values[0]:
                        continue
                    record = {"gene_id_version": values[0], "gene_id": values[0].split(".")[0],
                              "log2_fold_change": values[1], "p_value": values[2], "fdr": values[3],
                              "gene_name": values[4], "source_chromosome": str(values[5]),
                              "source_start": values[6], "source_end": values[7],
                              "source": {"source_id": name, "sheet": sheet.title, "row": rowno,
                                         "sha256": source["sha256"]},
                              "source_overlap_flags": dict(zip(headers[8:], values[8:]))}
                    records.append(record)
                expression[context][label] = records
    write_json("workbook_inventory.json", inventories)
    write_json("expression.json", expression)
    return sources, inventories, expression


def candidates(inventories: list[dict]) -> list[dict]:
    # Verify named intervals against two independent locations in source pack:
    # Supplement 5 sheet titles and Supplement 1 coordinate/width table.
    supp1 = next(x for x in inventories if x["source_id"].endswith("supp1.xlsx"))
    supp5 = next(x for x in inventories if x["source_id"].endswith("supp5.xlsx"))
    table = next(x for x in supp1["sheets"] if x["sheet"] == "overlap safe-active")
    result = []
    for sheet in supp5["sheets"][1:4]:
        prefix = next(x for x in NAMES if sheet["sheet"].startswith(x))
        candidate_id, name, chrom = NAMES[prefix]
        start, end = map(int, re.search(r"\((\d+)-(\d+)\)", sheet["sheet"]).groups())
        matches = [x for x in table["rows"] if x["values"][1:3] == [start, end]]
        assert len(matches) == 1, (name, matches)
        row = matches[0]
        assert str(row["values"][0]) == chrom[3:]
        assert row["values"][3] == end - start + 1
        result.append({"candidate_id": candidate_id, "name": name, "assembly": "GRCh38", "chromosome": chrom,
                       "start": start - 1, "end": end, "coordinate_system": "zero_based_half_open",
                       "cell_context": "H1 human embryonic stem cells", "data_version": VERSION,
                       "group_id": f"autio-{chrom}-{row['values'][4]}-{row['values'][5]}",
                       "source_coordinates": {"assembly": "GRCh38", "chromosome": chrom[3:], "start": start,
                                              "end": end, "width": end-start+1, "coordinate_system": "one_based_closed",
                                              "conversion": "internal_start = source_start - 1; internal_end = source_end",
                                              "basis": "Supplement 1 width=end-start+1; matched Supplement 5 sheet title; study GRCh38 methods",
                                              "source_id": supp1["source_id"], "sheet": table["sheet"], "row": row["source_row"],
                                              "sha256": supp1["sha256"]},
                       "evidence_ids": ["source:supp1", f"source:supp5:{candidate_id}"],
                       "origin": "publication-derived candidate shortlist", "publication_blat_ratio": row["values"][13]})
    write_json("candidates.json", result)
    return result


def references(candidate_records: list[dict]) -> tuple[list[dict], dict]:
    sources = []
    for name in ("wgEncodeGencodeCompV36.sql", "wgEncodeGencodeAttrsV36.sql", "wgEncodeGencodePseudoGeneV36.sql",
                 "wgEncodeGencodeCompV36.txt.gz", "wgEncodeGencodeAttrsV36.txt.gz", "wgEncodeGencodePseudoGeneV36.txt.gz", "chromInfo.txt.gz"):
        source = download("https://hgdownload.soe.ucsc.edu/goldenPath/hg38/database/" + name, name)
        source.update(attribution="UCSC Genome Browser; GENCODE v36 (where applicable)",
                      license_url="https://www.gencodegenes.org/pages/data_access.html")
        sources.append(source)
    attrs = {}
    with gzip.open(RAW / "wgEncodeGencodeAttrsV36.txt.gz", "rt") as fh:
        for line in fh:
            fields = line.rstrip("\n").split("\t")
            assert len(fields) == 14
            attrs[fields[4]] = {"gene_id_version": fields[0], "gene_id": fields[0].split(".")[0],
                                "name": fields[1], "gene_type": fields[2]}
    chromosomes = []
    canonical = [f"chr{i}" for i in range(1, 23)] + ["chrX", "chrY", "chrM"]
    with gzip.open(RAW / "chromInfo.txt.gz", "rt") as fh:
        for line in fh:
            chrom, length, *_ = line.rstrip().split("\t")
            if chrom in canonical:
                chromosomes.append({"chromosome": chrom, "length": int(length)})
    chromosomes.sort(key=lambda r: canonical.index(r["chromosome"]))
    write_json("chromosomes.json", chromosomes)
    transcripts, genes = [], {}
    for table_name in ("wgEncodeGencodeCompV36.txt.gz", "wgEncodeGencodePseudoGeneV36.txt.gz"):
        with gzip.open(RAW / table_name, "rt") as fh:
            for rowno, line in enumerate(fh, 1):
                f = line.rstrip("\n").split("\t")
                assert len(f) == 16
                if f[2] not in canonical:
                    continue
                attr = attrs[f[1]]
                item = {**attr, "transcript_id": f[1], "chromosome": f[2], "strand": f[3],
                        "start": int(f[4]), "end": int(f[5]), "source_row": rowno, "source_table": table_name,
                        "tss": int(f[4]) if f[3] == "+" else int(f[5])-1}
                key = (item["gene_id"], item["chromosome"])
                if key not in genes:
                    genes[key] = {k: v for k, v in item.items() if k not in ("transcript_id", "tss")}
                else:
                    genes[key]["start"] = min(genes[key]["start"], item["start"])
                    genes[key]["end"] = max(genes[key]["end"], item["end"])
                if any(item["chromosome"] == c["chromosome"] and item["end"] > c["start"]-1_000_000 and item["start"] < c["end"]+1_000_000 for c in candidate_records):
                    item["exons"] = list(zip(map(int, f[9].rstrip(",").split(",")), map(int, f[10].rstrip(",").split(","))))
                    transcripts.append(item)
    write_json("genes.json", list(genes.values()))
    assets = {}
    for c in candidate_records:
        start, end = c["start"]-100, c["end"]+100
        name = f"{c['candidate_id']}-hg38-sequence.json"
        url = f"https://api.genome.ucsc.edu/getData/sequence?genome=hg38;chrom={c['chromosome']};start={start};end={end}"
        source = download(url, name)
        sources.append(source)
        response = json.loads((RAW / name).read_text())
        sequence = response["dna"]
        assert len(sequence) == end-start and response["start"] == start and response["end"] == end
        seq = {"assembly": "GRCh38", "chromosome": c["chromosome"], "start": start, "end": end,
               "coordinate_system": "zero_based_half_open", "sequence": sequence,
               "orientation": "reference_plus", "sha256": digest(sequence.encode()), "source_hash": source["sha256"],
               "source_url": url, "personalized_h1_genome": False}
        features = [x for x in transcripts if x["chromosome"] == c["chromosome"] and x["end"] > c["start"]-1_000_000 and x["start"] < c["end"]+1_000_000]
        assets[c["candidate_id"]] = {"sequence": seq, "annotation": {"features": features,
                     "coverage": {"chromosome": c["chromosome"], "start": c["start"]-1_000_000, "end": c["end"]+1_000_000, "complete": True,
                                  "tables": ["wgEncodeGencodeCompV36", "wgEncodeGencodePseudoGeneV36"],
                                  "completeness_definition": "All overlapping rows from both frozen UCSC v36 transcript tables; canonical chromosomes only."},
                     "release": "GENCODE v36", "assembly": "GRCh38", "coordinate_system": "zero_based_half_open",
                     "source_hashes": {x["source_id"]: x["sha256"] for x in sources if "Gencode" in x["source_id"]},
                     "analysis_note": "New GENCODE v36 analysis; not an exact reproduction of the publication annotation pipeline."}}
        time.sleep(1)
    write_json("reference_assets.json", assets)
    return sources, assets


def main():
    RAW.mkdir(parents=True, exist_ok=True)
    sources, inventory, expression = workbooks()
    candidate_records = candidates(inventory)
    ref_sources, assets = references(candidate_records)
    sources += ref_sources
    manifest = {"schema_version": 1, "data_version": VERSION, "mode": "real_data", "assembly": "GRCh38",
                "criteria_sha256": digest((NORMALIZED / "criteria.json").read_bytes()),
                "cell_context": "H1 human embryonic stem cells", "study_url": STUDY, "sources": sources,
                "input_contract": {"expression_tables": "significant-only DEseq2 results, not raw counts or all measured genes",
                                   "gene_unit": "unique versioned Ensembl gene identifier; version-stripped counts also reported",
                                   "thresholds": {"abs_log2_fold_change_gte": 1, "fdr_lte": 0.01},
                                   "comparison": "targeted clone or independent untargeted culture versus reference untargeted cells of same H1/H9 context",
                                   "h9_policy": "Separate biological context; no automatic generalization to H1",
                                   "interpretive_conclusions": "Excluded from worker inputs; only tables, methods metadata and annotations are provided"},
                "unavailable": ["COSMIC cancer-gene data/licensing", "DHS exclusion track", "ultraconserved interval track", "H1 cCRE child asset"]}
    write_json("manifest.json", manifest)
    print(json.dumps({"candidates": len(candidate_records), "sources": len(sources),
                      "H1_table_counts": {k: len(v) for k,v in expression['H1'].items()},
                      "sequence_lengths": {k: len(v['sequence']['sequence']) for k,v in assets.items()}}, indent=2))


if __name__ == "__main__":
    main()
