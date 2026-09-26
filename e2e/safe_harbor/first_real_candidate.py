#!/usr/bin/env python3
"""SH-Q01 — first real candidate end to end.

Real OpenRouter model → approved scientific tools → MongoDB acceptance → export →
independent recomputation from the original hashed workbook → actual browser.

Run from the repository root with OPENROUTER_API_KEY and MODEL_ID in the ignored .env:
  PYTHONPATH=backend:. .venv/bin/python e2e/safe_harbor/first_real_candidate.py

Every attempted run is reported, including failed/blocked ones. Nothing is retried
silently and no assertion is relaxed to manufacture a pass.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[2]
load_dotenv(ROOT / ".env", override=False)

from safe_harbor.runtime.ledger import digest, now  # noqa: E402
from safe_harbor.science import run_tool  # noqa: E402

TERMINAL = {"complete", "blocked", "failed", "stopped", "budget_exhausted"}
SAFE_LABEL = re.compile(r"\b(is|are|as)\s+(a\s+)?(globally\s+)?safe\b(?!\s*(or|/))", re.I)


def request(base: str, path: str, body=None, timeout=10):
    data = None if body is None else json.dumps(body).encode()
    req = urllib.request.Request(base + path, data=data, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return json.load(response)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def flatten_numbers(value, out: set[float]):
    if isinstance(value, bool):
        return
    if isinstance(value, (int, float)):
        out.add(float(value))
    elif isinstance(value, dict):
        for item in value.values():
            flatten_numbers(item, out)
    elif isinstance(value, list):
        for item in value:
            flatten_numbers(item, out)


def independent_expression(candidate_sheet_prefix: str) -> dict:
    """Recount from the original supplementary workbook, not from normalized JSON or tool code."""
    import openpyxl
    workbook = openpyxl.load_workbook(ROOT / "data/safe_harbor/raw/elife-79592-supp5.xlsx", read_only=True)

    def significant(sheet_name: str) -> dict[str, float]:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        header = rows[0]
        lfc = next(i for i, name in enumerate(header) if str(name).endswith("_LFC"))
        fdr = next(i for i, name in enumerate(header) if str(name).endswith("_FDR"))
        genes = {}
        for row in rows[1:]:
            if row[0] and row[lfc] is not None and row[fdr] is not None and abs(row[lfc]) >= 1 and row[fdr] <= 0.01:
                genes[row[0]] = row[lfc]
        return genes

    sheet = next(name for name in workbook.sheetnames if name.startswith(candidate_sheet_prefix + " ("))
    targeted, control = significant(sheet), significant("WT H1")
    shared = set(targeted) & set(control)
    same = [gene for gene in shared if (targeted[gene] > 0) == (control[gene] > 0)]
    return {
        "sheet": sheet, "targeted": len(targeted), "targeted_up": sum(v > 0 for v in targeted.values()),
        "targeted_down": sum(v < 0 for v in targeted.values()), "control": len(control),
        "control_up": sum(v > 0 for v in control.values()), "control_down": sum(v < 0 for v in control.values()),
        "shared": len(shared), "shared_same_direction": len(same), "shared_opposite_direction": len(shared) - len(same),
        "shared_percent_of_targeted": round(100 * len(shared) / len(targeted), 2) if targeted else None,
    }


class Journey:
    def __init__(self, args):
        self.args = args
        self.base = f"http://127.0.0.1:{args.api_port}"
        stamp = time.strftime("%Y%m%dT%H%M%S")
        self.output = ROOT / "artifacts/safe_harbor/first-real-candidate" / stamp
        self.output.mkdir(parents=True, exist_ok=True)
        self.database = f"safe_harbor_q01_{stamp}"
        self.processes: list[subprocess.Popen] = []
        self.checks: list[dict] = []
        self.report = {
            "ticket": "SH-Q01", "started_at": now(), "mode": "real_model", "database": self.database,
            "model_id": os.getenv("MODEL_ID"), "model_provider": "openrouter", "candidate_id": args.candidate,
            "budget_request": {"token_limit": args.token_limit} if args.token_limit else None,
            "attempts": [], "checks": self.checks, "findings": [],
        }

    def check(self, name: str, ok: bool, evidence):
        self.checks.append({"check": name, "passed": bool(ok), "evidence": evidence})
        print(("PASS " if ok else "FAIL ") + name, flush=True)

    def start(self, command: list[str], env: dict, log_name: str, cwd: Path = ROOT):
        log = open(self.output / log_name, "w")
        process = subprocess.Popen(command, cwd=cwd, env=env, stdout=log, stderr=subprocess.STDOUT)
        self.processes.append(process)
        return process

    def stop(self):
        for process in reversed(self.processes):
            if process.poll() is None:
                process.terminate()
                try:
                    process.wait(10)
                except subprocess.TimeoutExpired:
                    process.kill()

    def start_api(self):
        env = {**os.environ, "MONGODB_DATABASE": self.database, "PYTHONPATH": f"{ROOT / 'backend'}:{ROOT}"}
        self.start([sys.executable, "-m", "uvicorn", "safe_harbor.api:app", "--host", "127.0.0.1", "--port", str(self.args.api_port)], env, "api.log")
        for _ in range(60):
            try:
                health = request(self.base, "/health")
                break
            except Exception:
                time.sleep(0.5)
        else:
            raise RuntimeError("API did not become healthy")
        self.check("api healthy with MongoDB authority and model configured", health["authoritative_store"] == "MongoDB" and health["model_available"], health)

    def attempt_run(self) -> tuple[str, dict] | tuple[None, None]:
        body = {"candidate_ids": [self.args.candidate], "mode": "real_model"}
        if self.args.token_limit:
            body["budget"] = {"token_limit": self.args.token_limit}
        for index in range(self.args.attempts):
            run_id = request(self.base, "/runs", body)["run_id"]
            started = time.monotonic()
            while True:
                snapshot = request(self.base, f"/runs/{run_id}/snapshot")
                if snapshot["run"]["status"] in TERMINAL or time.monotonic() - started > self.args.run_timeout:
                    break
                time.sleep(2)
            run = snapshot["run"]
            self.report["attempts"].append({
                "attempt": index + 1, "run_id": run_id, "status": run["status"], "duration_seconds": round(time.monotonic() - started, 1),
                "budget": run["budget"], "stop_reason": run.get("stop_reason"),
                "tasks": [{"kind": t["kind"], "status": t["status"], "error": t.get("error")} for t in snapshot["tasks"]],
            })
            print(f"attempt {index + 1}: {run_id} {run['status']}", flush=True)
            if run["status"] == "complete":
                return run_id, snapshot
        return None, None

    def verify(self, run_id: str):
        export = request(self.base, f"/runs/{run_id}/export", timeout=30)
        (self.output / "export.json").write_text(json.dumps(export, indent=1))
        snapshot = export["snapshot"]
        run, artifacts, tasks = snapshot["run"], snapshot["artifacts"], snapshot["tasks"]
        tool_artifacts = [a for a in artifacts if a["kind"] == "scientific_tool_result"]
        traces = [a for a in artifacts if a["kind"] == "worker_trace"]
        assessment = max(snapshot["assessments"], key=lambda a: a["assessment_revision"])

        self.check("run recorded as real_model with configured model identity", run["mode"] == "real_model" and run["model_id"] == os.getenv("MODEL_ID") and run["model_provider"] == "openrouter",
                   {k: run.get(k) for k in ("mode", "model_id", "model_provider", "harness_hash", "data_version")})
        usage = [t["data"]["usage"] for t in traces]
        self.check("every worker made real recorded model calls with usage", bool(usage) and all(u["model_calls"] >= 1 and u["tokens"] > 0 for u in usage),
                   {"workers": len(usage), "model_calls": sum(u["model_calls"] for u in usage), "tokens": sum(u["tokens"] for u in usage),
                    "cost_usd": run["budget"].get("cost_usd"), "uncertain_tokens": run["budget"].get("uncertain_tokens")})
        self.check("all tasks accepted", all(t["status"] == "complete" for t in tasks), [(t["kind"], t["status"]) for t in tasks])
        self.check("events are ordered and gap-free", [e["sequence"] for e in export["events"]] == list(range(1, len(export["events"]) + 1)), {"events": len(export["events"])})

        manifest = json.loads((ROOT / "data/safe_harbor/normalized/manifest.json").read_text())
        raw_paths = {Path(source["path"]).name: ROOT / source["path"] for source in manifest["sources"]}
        cited = {name: value for a in tool_artifacts for name, value in a["data"].get("source_hashes", {}).items()}
        mismatched = {name: value for name, value in cited.items() if name not in raw_paths or not raw_paths[name].exists() or sha256(raw_paths[name]) != value}
        # UCSC sequence responses embed downloadTime; a re-fetched cache changes the body hash
        # without changing bases (see PROVENANCE.md). Those are verified on the bases instead,
        # and reported separately rather than hidden.
        assets = json.loads((ROOT / "data/safe_harbor/normalized/reference_assets.json").read_text())
        base_verified = {}
        for name in list(mismatched):
            asset = next((a["sequence"] for a in assets.values() if a["sequence"].get("source_hash") == mismatched[name]), None)
            if asset and raw_paths.get(name, Path("/nonexistent")).exists():
                bases = json.loads(raw_paths[name].read_text())["dna"]
                if hashlib.sha256(bases.encode()).hexdigest() == asset["sha256"]:
                    base_verified[name] = {"body_sha256_on_disk": sha256(raw_paths[name]), "committed_body_sha256": mismatched.pop(name), "bases_sha256": asset["sha256"]}
        self.check("every cited source matches original bytes (sequence JSON: identical bases)", not mismatched and bool(cited), {"cited": cited, "mismatched": mismatched, "wrapper_metadata_drift_bases_identical": base_verified})

        recomputed = []
        for trace in traces:
            for call in trace["data"]["tool_calls"]:
                fresh = run_tool(call["tool_name"], self.args.candidate, call["arguments"])
                recomputed.append({"tool": call["tool_name"], "arguments": call["arguments"], "stored": digest(call["result"]["calculation"]), "recomputed": digest(fresh["calculation"])})
        self.check("every model-requested tool result recomputes identically from exported arguments", bool(recomputed) and all(r["stored"] == r["recomputed"] for r in recomputed), recomputed)

        sheet_prefix = {"pansio-1": "Pansio-1", "olonne-18": "Olônne-18", "keppel-19": "Keppel-19"}[self.args.candidate]
        independent = independent_expression(sheet_prefix)
        comparison = next((a["data"]["calculation"] for a in tool_artifacts if a["data"]["tool_name"] == "expression_comparison"), None)
        overlap = next((a["data"]["calculation"] for a in tool_artifacts if a["data"]["tool_name"] == "control_overlap"), None)
        tool_numbers: set[float] = set()
        flatten_numbers(comparison, tool_numbers)
        flatten_numbers(overlap, tool_numbers)
        expected = {k: independent[k] for k in ("targeted", "targeted_up", "targeted_down", "control", "shared", "shared_same_direction")}
        self.check("independent workbook recount agrees with the tool calculations", comparison is not None and overlap is not None and all(float(v) in tool_numbers for v in expected.values()),
                   {"independent_from_raw_workbook": independent, "tool_expression_comparison_targeted": (comparison or {}).get("targeted")})

        text_fields = [assessment.get("conclusion", "")] + [r.get("reason", "") for r in assessment["experimental_endpoint_results"]]
        text = " ".join(text_fields)
        run_numbers: set[float] = set()
        for a in tool_artifacts:
            flatten_numbers(a["data"]["calculation"], run_numbers)
        run_numbers |= {float(v) for v in independent.values() if isinstance(v, (int, float))}
        run_numbers |= {float(len(assessment["criterion_results"])), float(sum(r["status"] == "pass" for r in assessment["criterion_results"])), float(sum(r["status"] != "pass" for r in assessment["criterion_results"]))}
        criteria_text = (ROOT / "data/safe_harbor/normalized/criteria.json").read_text()
        for source_text in [criteria_text] + assessment["limitations"] + [l for a in tool_artifacts for l in a["data"].get("limitations", [])] + [r.get("reason", "") for r in assessment["criterion_results"]]:
            # Version numbers and thresholds quoted from frozen criteria or the run's recorded limitations.
            run_numbers |= {float(m.replace(",", "")) for m in re.findall(r"\d[\d,]*(?:\.\d+)?", source_text)}
        quoted = {float(m.replace(",", "")) for m in re.findall(r"(?<![\w.-])(\d[\d,]*(?:\.\d+)?)(?![\w.])", text)}
        unsupported = sorted(q for q in quoted if q >= 10 and not any(abs(q - n) < 0.006 for n in run_numbers))
        self.check("every multi-digit number in the conclusion is backed by a run calculation", not unsupported, {"quoted": sorted(quoted), "unsupported": unsupported})

        self.check("assessment keeps independent axes and required fields", assessment["screen_status"] in {"pass", "fail", "incomplete"} and assessment["evidence_status"] in {"supported_for_endpoint", "conflicting", "unknown"} and assessment["freshness"] in {"current", "stale"} and all(k in assessment for k in ("candidate_id", "assembly", "cell_context", "criterion_results", "experimental_endpoint_results", "evidence_ids", "unresolved_questions", "limitations", "assessment_revision")),
                   {k: assessment[k] for k in ("screen_status", "evidence_status", "freshness", "assessment_revision", "assembly", "cell_context")})
        unavailable = [r["criterion_id"] for r in assessment["criterion_results"] if r["status"] != "pass"]
        self.check("missing required evidence yields incomplete, not pass", not unavailable or assessment["screen_status"] != "pass", {"non_pass_criteria": unavailable, "screen_status": assessment["screen_status"]})
        self.check("no global safe label in accepted text", not SAFE_LABEL.search(text), {"matches": [m.group(0) for m in SAFE_LABEL.finditer(text)]})
        known_ids = {a["artifact_id"] for a in artifacts} | {e for a in tool_artifacts for e in a["data"].get("evidence_ids", [])}
        dangling = sorted(set(assessment["evidence_ids"]) - known_ids)
        self.check("every assessment evidence ID resolves to a run artifact or tool evidence", not dangling, {"evidence_ids": len(assessment["evidence_ids"]), "dangling": dangling})

        conclusion = assessment.get("conclusion", "")
        typed = True
        try:
            parsed = json.loads(conclusion.strip().removeprefix("```json").removesuffix("```"))
            typed = not isinstance(parsed, dict)
        except (json.JSONDecodeError, AttributeError):
            pass
        self.check("accepted conclusion is typed prose, not an unparsed model JSON envelope", typed and not conclusion.lstrip().startswith("{"), {"conclusion_prefix": conclusion[:160]})
        return assessment

    def browser(self, run_id: str, assessment: dict):
        env = {**os.environ, "API_TARGET": self.base}
        self.start(["npx", "vite", "--host", "127.0.0.1", "--port", str(self.args.ui_port), "--strictPort"], env, "vite.log", cwd=ROOT / "frontend")
        ui = f"http://127.0.0.1:{self.args.ui_port}"
        for _ in range(60):
            try:
                urllib.request.urlopen(ui, timeout=2)
                break
            except Exception:
                time.sleep(0.5)
        expected = {"run_id": run_id, "candidate_id": self.args.candidate, "conclusion": assessment["conclusion"], "screen_status": assessment["screen_status"],
                    "evidence_status": assessment["evidence_status"], "revision": assessment["assessment_revision"]}
        result = subprocess.run(["node", str(ROOT / "e2e/safe_harbor/first_real_candidate_browser.mjs"), ui, json.dumps(expected), str(self.output)], capture_output=True, text=True, timeout=180)
        try:
            observed = json.loads(result.stdout.strip().splitlines()[-1])
        except (IndexError, json.JSONDecodeError):
            observed = {"error": (result.stdout + result.stderr)[-2000:]}
        for check in observed.get("checks", []):
            self.check("browser: " + check["check"], check["passed"], check.get("evidence"))
        if "error" in observed:
            self.check("browser journey completed", False, observed)

    def run(self):
        try:
            self.start_api()
            run_id, _ = self.attempt_run()
            self.check("a real-model run reached complete within the attempt limit", run_id is not None, {"attempts": len(self.report["attempts"])})
            if run_id:
                self.report["run_id"] = run_id
                assessment = self.verify(run_id)
                self.browser(run_id, assessment)
        except Exception as exc:  # recorded, never swallowed into a pass
            self.check("journey executed without exception", False, repr(exc))
        finally:
            self.stop()
            self.report["finished_at"] = now()
            self.report["passed"] = bool(self.checks) and all(c["passed"] for c in self.checks)
            (self.output / "report.json").write_text(json.dumps(self.report, indent=1))
            print(json.dumps({"passed": self.report["passed"], "report": str(self.output / "report.json")}))
        return 0 if self.report["passed"] else 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--candidate", default="pansio-1", choices=["pansio-1", "olonne-18", "keppel-19"])
    parser.add_argument("--attempts", type=int, default=3)
    parser.add_argument("--token-limit", type=int, default=None, help="Explicit, reported run token budget override")
    parser.add_argument("--run-timeout", type=int, default=600)
    parser.add_argument("--api-port", type=int, default=8050)
    parser.add_argument("--ui-port", type=int, default=5195)
    args = parser.parse_args()
    if not (os.getenv("OPENROUTER_API_KEY") and os.getenv("MODEL_ID")):
        sys.exit("OPENROUTER_API_KEY and MODEL_ID must be configured in the ignored .env file")
    sys.exit(Journey(args).run())


if __name__ == "__main__":
    main()
